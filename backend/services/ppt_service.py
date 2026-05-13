from pptx import Presentation
from pptx.util import Inches
import os
from fastapi import Header, HTTPException

REPORT_DIR = "reports"

def generate_ppt(file_id: str, charts: list):
    prs = Presentation()
    os.makedirs(REPORT_DIR, exist_ok=True)

    # Title slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "AI Analyst Report"
    slide.placeholders[1].text = f"File ID: {file_id}"

    # Add chart slides
    for chart in charts:
        BASE_DIR = os.path.dirname(os.path.abspath(__file__))
        PROJECT_ROOT = os.path.abspath(os.path.join(BASE_DIR, ".."))
        
        if os.path.isabs(chart):
            chart_full_path = chart
        else:
            chart_full_path = os.path.join(PROJECT_ROOT, chart)
        
        if not os.path.exists(chart_full_path):
            continue

        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = os.path.basename(chart)

        slide.shapes.add_picture(
            chart_full_path,
            Inches(1),
            Inches(1.5),
            width=Inches(6)
        )

    ppt_path = f"{REPORT_DIR}/{file_id}_report.pptx"
    prs.save(ppt_path)

    return ppt_path

INTERNAL_API_KEY = "internal-secret"


def verify_internal_key(
    x_internal_key: str = Header(None)
):

    if x_internal_key != INTERNAL_API_KEY:

        raise HTTPException(
            status_code=401,
            detail="Unauthorized"
        )
