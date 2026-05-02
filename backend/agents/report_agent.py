import os
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet
from pptx import Presentation

CLEAN_DIR = "cleaned"
CHART_DIR = "charts"
REPORT_DIR = "reports"

os.makedirs(REPORT_DIR, exist_ok=True)


def generate_report(file_id: str, analysis: dict, charts: list):
    report_path = f"{REPORT_DIR}/{file_id}_report.pdf"

    doc = SimpleDocTemplate(report_path)
    styles = getSampleStyleSheet()

    content = []

    # Title
    content.append(Paragraph("AI Analyst Report", styles["Title"]))
    content.append(Spacer(1, 10))

    # Summary
    content.append(Paragraph("Executive Summary", styles["Heading2"]))
    content.append(Spacer(1, 8))

    summary_text = f"""
    Rows analyzed: {analysis.get('row_count')}<br/>
    Columns: {analysis.get('column_count')}
    """
    content.append(Paragraph(summary_text, styles["Normal"]))
    content.append(Spacer(1, 12))

    # Numeric Summary
    content.append(Paragraph("Key Metrics", styles["Heading2"]))
    content.append(Spacer(1, 8))

    for col, stats in analysis.get("numeric_summary", {}).items():
        text = f"{col} → Mean: {stats['mean']}, Min: {stats['min']}, Max: {stats['max']}"
        content.append(Paragraph(text, styles["Normal"]))

    content.append(Spacer(1, 12))

    # Charts
    content.append(Paragraph("Visualizations", styles["Heading2"]))
    content.append(Spacer(1, 10))

    for chart in charts:
        if os.path.exists(chart):
            content.append(Image(chart, width=400, height=250))
            content.append(Spacer(1, 10))

    doc.build(content)

    ppt_path = generate_ppt(file_id, analysis, charts)

    return {
        "report_path": report_path,
        "ppt_path": ppt_path
    }


def generate_ppt(file_id, analysis, charts):
    ppt_path = f"reports/{file_id}_report.pptx"

    prs = Presentation()

    slide_layout = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide_layout.shapes.title
    title.text = "AI Analyst Report"

    prs.save(ppt_path)

    return ppt_path